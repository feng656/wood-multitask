from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


INLINE_MARKERS = re.compile(r"(\*\*|`)")


def clean_inline(text: str) -> str:
    return INLINE_MARKERS.sub("", text.strip())


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    heading_sizes = {"Title": 20, "Heading 1": 15, "Heading 2": 13}
    for name, size in heading_sizes.items():
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = None
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("木材年轮-缺陷多任务项目")


def add_table(document: Document, lines: list[str]) -> None:
    rows = [
        [clean_inline(cell) for cell in line.strip().strip("|").split("|")]
        for line in lines
    ]
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    if not rows:
        return

    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index in range(width):
            cell = table.cell(row_index, column_index)
            value = row[column_index] if column_index < len(row) else ""
            cell.text = value
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(
                        qn("w:eastAsia"), "宋体"
                    )
                    if row_index == 0:
                        run.bold = True
            if row_index == 0:
                set_cell_shading(cell, "D9EAF2")


def export_markdown(source: Path, destination: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)

    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph()
                paragraph.style = document.styles["Normal"]
                paragraph.paragraph_format.left_indent = Cm(0.6)
                paragraph.paragraph_format.space_after = Pt(6)
                run = paragraph.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                code_lines.clear()
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while index < len(lines):
                candidate = lines[index].strip()
                if not (candidate.startswith("|") and candidate.endswith("|")):
                    break
                table_lines.append(candidate)
                index += 1
            add_table(document, table_lines)
            continue

        if stripped.startswith("# "):
            paragraph = document.add_paragraph(clean_inline(stripped[2:]), style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith("## "):
            document.add_heading(clean_inline(stripped[3:]), level=1)
        elif stripped.startswith("### "):
            document.add_heading(clean_inline(stripped[4:]), level=2)
        elif stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(clean_inline(stripped[2:]))
        elif stripped:
            paragraph = document.add_paragraph(clean_inline(stripped))
            paragraph.paragraph_format.space_after = Pt(6)

        index += 1

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export project Markdown reports to DOCX.")
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    export_markdown(args.source, args.destination)

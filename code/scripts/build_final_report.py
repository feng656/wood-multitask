from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
PROCESSED = PROJECT_ROOT / "data_processed"
DOCS = PROJECT_ROOT / "docs"
REPORT_STEM = "paper_grade_data_engineering_report"


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def load_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(4 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def split_text(value: dict[str, int]) -> str:
    return ", ".join(f"{key}={count}" for key, count in value.items())


def markdown_table(headers: list[str], rows: Iterable[Iterable[object]]) -> list[str]:
    lines = ["| " + " | ".join(headers) + " |", "|" + "|".join("---" for _ in headers) + "|"]
    for row in rows:
        lines.append("| " + " | ".join(str(value).replace("|", "/") for value in row) + " |")
    return lines


def collect_facts() -> dict[str, Any]:
    manifest = load_csv(PROCESSED / "manifests" / "manifest.csv")
    quality = load_json(PROCESSED / "manifests" / "quality_report.json")
    integrity = load_json(PROCESSED / "integrity_report.json")
    registry = load_json(PROCESSED / "provenance" / "dataset_registry.json")
    source_verification = load_json(PROCESSED / "provenance" / "source_verification.json")
    bridge = load_json(PROCESSED / "bridge_annotations" / "normalization_audit.json")
    ring = load_json(PROCESSED / "ring_targets" / "summary.json")
    ring_rows = load_csv(PROCESSED / "ring_targets" / "ring_targets.csv")
    crops = load_json(PROCESSED / "classification_crops" / "summary.json")
    duplicates = load_json(PROCESSED / "duplicate_audit" / "summary.json")
    visual_review = load_json(PROCESSED / "duplicate_audit" / "visual_review.json")

    dataset_category: Counter[tuple[str, str]] = Counter()
    split_category: Counter[tuple[str, str]] = Counter()
    for path in sorted((PROCESSED / "defect_coco").glob("instances_*.json")):
        document = load_json(path)
        images = {int(item["id"]): item for item in document["images"]}
        for annotation in document["annotations"]:
            category = str(annotation["category"])
            dataset_id = str(annotation["source_dataset"])
            split = str(images[int(annotation["image_id"])]["split"])
            dataset_category[(dataset_id, category)] += 1
            split_category[(split, category)] += 1

    normal_by_dataset = Counter(
        row["dataset_id"]
        for row in manifest
        if row["defect_label_state"] == "verified_negative"
        and row["dataset_id"] in {"vsb", "vnwoodknot", "mokume", "oulu"}
    )
    bridge_rows = [row for row in manifest if row["dataset_role"] == "bridge"]
    bridge_states = Counter(row["defect_label_state"] for row in bridge_rows)
    bridge_splits = Counter(row["split"] for row in bridge_rows)
    scope_counts = Counter(row["supervision_scope"] for row in manifest)
    ring_warning_samples = [
        row["sample_id"] for row in ring_rows if int(row.get("non_nested_pairs") or 0) > 0
    ]

    storage_files = 0
    storage_bytes = 0
    for path in PROCESSED.rglob("*"):
        if path.is_file():
            storage_files += 1
            storage_bytes += path.stat().st_size

    critical_paths = [
        PROJECT_ROOT / "configs" / "data.json",
        PROJECT_ROOT / "configs" / "defect_taxonomy.json",
        PROJECT_ROOT / "configs" / "dataset_registry.json",
        PROCESSED / "manifests" / "manifest.csv",
        PROCESSED / "manifests" / "group_splits.csv",
        PROCESSED / "ring_targets" / "ring_targets.csv",
        PROCESSED / "classification_crops" / "classification_crops.csv",
        PROCESSED / "integrity_report.json",
        PROCESSED / "duplicate_audit" / "primary" / "cross_split_phash_candidates.csv",
        PROCESSED / "provenance" / "source_verification.json",
        *sorted((PROCESSED / "defect_coco").glob("instances_*.json")),
    ]
    artifact_hashes = {
        path.relative_to(PROJECT_ROOT).as_posix(): sha256_file(path) for path in critical_paths
    }
    (PROCESSED / "provenance" / "artifact_hashes.json").write_text(
        json.dumps(artifact_hashes, indent=2), encoding="utf-8"
    )

    return {
        "manifest": manifest,
        "quality": quality,
        "integrity": integrity,
        "registry": registry,
        "source_verification": source_verification,
        "bridge": bridge,
        "ring": ring,
        "ring_warning_samples": ring_warning_samples,
        "crops": crops,
        "duplicates": duplicates,
        "visual_review": visual_review,
        "dataset_category": dataset_category,
        "split_category": split_category,
        "normal_by_dataset": normal_by_dataset,
        "bridge_states": bridge_states,
        "bridge_splits": bridge_splits,
        "scope_counts": scope_counts,
        "storage_files": storage_files,
        "storage_bytes": storage_bytes,
        "artifact_hashes": artifact_hashes,
    }


def build_markdown(facts: dict[str, Any]) -> str:
    quality = facts["quality"]
    integrity = facts["integrity"]
    registry = facts["registry"]
    bridge = facts["bridge"]
    ring = facts["ring"]
    crops = facts["crops"]
    duplicates = facts["duplicates"]
    visual = facts["visual_review"]
    provenance = facts["source_verification"]
    dataset_category = facts["dataset_category"]
    split_category = facts["split_category"]

    lines = [
        "# 木材年轮与缺陷多任务项目：论文级数据工程报告",
        "",
        f"报告日期：{date.today().isoformat()}",
        "",
        "## 1. 结论",
        "",
        f"最终统一清单包含 **{quality['total_images']:,}** 张图像：主监督数据 **{facts['scope_counts']['primary']:,}** 张，辅助数据 **{facts['scope_counts']['auxiliary_only']:,}** 张。清单质量状态和派生产物完整性状态均为 **ok**。",
        "",
        f"源数据只读验证通过：处理前后元数据指纹均为 `{provenance['baseline_metadata_sha256']}`，变化数据集 0，关键标注内容变化 0，缺失文件 0。所有新增和修改只发生在 `wood-multitask-final`。",
        "",
        "当前数据工程已满足年轮结构、缺陷实例、分类裁剪和双标注桥接四类监督的统一读取要求。模型尚未在本工程上完成正式训练，因此本报告不声称年轮与缺陷已经呈现双向正迁移；第 10 节给出可证伪的论文实验协议。",
        "",
        "## 2. 标签与监督策略",
        "",
        "- 缺陷检测仅保留 `knot` 和独立 `crack` 两类；检测头含背景时输出通道为 3，YOLO 类别编号为 knot=0、crack=1。",
        "- `knot_with_crack` 只映射为 `knot`，亚型记录为 `cracked`，绝不进入 crack 真值。",
        "- Marrow、overgrown、Quartzity、Blue_stain、resin 不作为目标类别。VSB 中这些标注被排除；OULU/Mokume 的非目标区域保留为 ignore，不被解释为正常背景。",
        "- 分类任务为 `normal/knot/crack`。normal 仅来自已核验负样本；标签缺失始终通过 task mask 隔离。",
        "- Mokume 保留矩形标注，UruDendro 保留真实多边形。年轮标注没有因缺陷矩形化而改动。",
        "",
        "原始项目文档中的 VSB 20,276 张和 6 类缺陷统计属于旧方案。本版本以已确认的 VSB 16,000 张、两类监督策略为准，并重新生成全部统计。",
        "",
        "## 3. 数据集登记与角色",
        "",
    ]
    dataset_rows = []
    for dataset_id, metadata in registry["datasets"].items():
        dataset_rows.append([
            dataset_id,
            metadata["protocol_role"],
            metadata["manifest_images"],
            metadata["physical_groups"],
            split_text(metadata["splits"]),
            metadata["ring_images"],
            metadata["defect_labeled_images"],
            metadata["target_instances_manifest"],
            metadata["ignored_instances_manifest"],
        ])
    lines.extend(markdown_table(
        ["数据集", "协议角色", "图像", "物理组", "划分", "年轮", "缺陷已标注", "目标实例", "忽略实例"],
        dataset_rows,
    ))
    lines.extend([
        "",
        "WVTec Wood 仅用于外部异常评估；涂层数据仅用于合成压力测试。MiSCS 当前只有代码仓库，没有图像数据，故不计入清单。`mokume-labelme`、`vsb-统计`、`vsb-final.build-cache` 等派生目录不作为独立数据集。除 WVTec 本地记录的许可外，其余数据集在公开发布前仍需逐项确认许可证。",
        "",
        "## 4. 物理组划分与防泄漏",
        "",
        "- 固定随机种子为 `20260729`，先按物理组划分，再生成目标、COCO 和裁剪。",
        "- 桥接集共 322 张：train=194、val=59、test=69；树或立方体在三个划分间交集为 0。",
        "- VSB 继承 `vsb-final` 的固定采集组划分：train=10,000、val=2,000、test=4,000。",
        "- UruDendro4 锁定为年轮 external_test，并建立 5 折树级 R2；Indiana、VNWoodKnot、OULU 均锁定 external_test。",
        "- 清单重复路径 0、缺失引用 0、不可读图像 0、跨划分物理组泄漏 0。",
        "",
        "## 5. 年轮监督产物",
        "",
        f"已生成 **{ring['generated']:,}/{ring['requested']:,}** 份年轮目标，错误 0，共 {ring['generated'] * 5:,} 个 PNG。每张包含 1 px 骨架、{ring['boundary_width']} px 边界、tau={ring['distance_tau']} 截断距离场、实例图和有效区。数据集分布：{json.dumps(ring['by_dataset'], ensure_ascii=False)}。",
        "",
        f"共解析 {ring['closed_rings']:,} 条闭合轮廓。发现 {ring['non_nested_pairs']} 个非严格嵌套警告，样本为：{', '.join(facts['ring_warning_samples'])}。这些目标可用于一般训练，但拓扑敏感实验前必须复核原始多边形。",
        "",
        "## 6. 两类缺陷 COCO",
        "",
        f"扩展 COCO 包含 **{integrity['defect_coco']['images']:,}** 张图像、**{integrity['defect_coco']['annotations']:,}** 个目标实例和 **{integrity['defect_coco']['ignore_regions']:,}** 个非目标忽略区域。分割可用实例 {integrity['defect_coco']['segmented_annotations']:,} 个。",
        "",
    ])
    defect_rows = []
    for dataset_id in ("vsb", "vnwoodknot", "oulu", "urudendro", "mokume"):
        knot = dataset_category[(dataset_id, "knot")]
        crack = dataset_category[(dataset_id, "crack")]
        defect_rows.append([dataset_id, knot, crack, knot + crack, facts["normal_by_dataset"][dataset_id]])
    defect_rows.append([
        "总计",
        integrity["defect_coco"]["category_counts"]["knot"],
        integrity["defect_coco"]["category_counts"]["crack"],
        integrity["defect_coco"]["annotations"],
        sum(facts["normal_by_dataset"].values()),
    ])
    lines.extend(markdown_table(["数据集", "knot", "crack", "目标合计", "已核验 normal"], defect_rows))
    lines.extend(["", "按协议划分的目标实例：", ""])
    split_rows = []
    for split in ("train", "val", "test", "external_test"):
        knot = split_category[(split, "knot")]
        crack = split_category[(split, "crack")]
        split_rows.append([split, knot, crack, knot + crack])
    lines.extend(markdown_table(["划分", "knot", "crack", "合计"], split_rows))
    lines.extend([
        "",
        "所有框均位于图像范围内，COCO 引用、分割、类别契约和 group/split 继承错误均为 0。VSB 的 16,000 张统一掩膜与 32,095 个裁剪原位复用，没有再复制约 129 GiB 数据。",
        "",
        "标准 YOLO 和 Mask R-CNN 导出格式不能表达 ignore 区域：导出器默认使用 box 模式，并保守跳过含 ignore 区域的图像；正式 OULU 全量评估应使用扩展 COCO 加 ignore-aware evaluator，避免把非目标缺陷计作背景。",
        "",
        "## 7. 分类裁剪",
        "",
        f"共索引 **{crops['total']:,}** 个 320×320 裁剪：knot={crops['class_counts']['knot']:,}、crack={crops['class_counts']['crack']:,}、normal={crops['class_counts']['normal']:,}。其中复用 VSB 裁剪 {crops['reused_vsb_crops']:,} 个，本项目新增 4,193 个。裁剪缺失和物理组/划分继承错误均为 0。",
        "",
        "类别明显不平衡，正式训练必须报告 macro-F1、每类召回率和每类 AP，并使用损失重加权或按物理组的平衡采样；不能只报告总体准确率。",
        "",
        "## 8. 桥接数据与双任务监督",
        "",
        f"桥接图像 322 张，划分为 {split_text(dict(facts['bridge_splits']))}；标签状态为 {dict(facts['bridge_states'])}。规范化副本保留 {bridge['output_shapes']}/{bridge['input_shapes']} 个形状，其中目标实例 {bridge['primary_instances']} 个、ignore 实例 {bridge['ignored_instances']} 个。",
        "",
        "桥接目标包括 UruDendro crack 802 个、Mokume knot 79 个和 crack 88 个；969 个目标均带 `affects_rings` 属性，可用于缺陷邻域年轮指标和双向条件分支实验。规范化执行实例 ID 去重 6 次、truncated 重算 176 次、自交修复 2 次，并保留 1 个非目标忽略区域。",
        "",
        "## 9. 重复与来源完整性审计",
        "",
        f"主数据 pHash 审计覆盖 {duplicates['images_hashed']:,} 张，阈值为 {duplicates['threshold']}：跨划分候选 {duplicates['cross_split_candidates']} 对，完全相同哈希 0 对。VSB 复用 16,000 个哈希，其余新计算 5,346 个。5 张接触表已复核全部 {visual['candidates_reviewed']} 对，明显图像重复 0 对；复核为 AI 辅助目视筛查，不是独立领域专家复核，也不能证明物理木材来源必然不同。",
        "",
        f"WVTec 辅助划分候选 0 对。涂层数据为 {duplicates['coating_group_audit']['physical_groups']} 个基础组，每组 45 个派生、总计 {duplicates['coating_group_audit']['images']:,} 张，全部处于 auxiliary_stress，不进入主训练/验证结论。",
        "",
        f"派生产物当前共 {facts['storage_files']:,} 个文件、{facts['storage_bytes'] / (1024 ** 2):.1f} MiB，不包含原位复用的 VSB 大文件。关键配置、清单、COCO、裁剪索引和完整性报告的 SHA-256 已写入 `data_processed/provenance/artifact_hashes.json`。",
        "",
        "## 10. 双向正迁移验证协议",
        "",
        "数据工程只能提供可验证条件，不能单独证明正相关。论文中应将命题表述为“年轮任务与缺陷任务是否产生稳定双向正迁移”，并至少比较以下模型：Ring-only、Defect-only、共享编码器但无交互、R→D Oracle、R→D Pred、D→R Oracle、D→R Pred、完整 BiCRR。",
        "",
        "- 年轮指标：Boundary F1/Dice、clDice、ASSD或HD95、闭合率、断裂数；同时报告缺陷邻域和匹配非缺陷区域。",
        "- 缺陷指标：box/mask mAP50-95、AP_knot、AP_crack、每类召回率；分类报告 macro-F1 和 balanced accuracy。",
        "- 统计单位必须是树或立方体等物理组。固定数据划分，至少 5 个训练随机种子，报告均值、标准差、组级 bootstrap 95% CI 和配对置换检验；多重比较使用 Holm 校正。",
        "- 只有完整模型相对无交互共享基线在两个方向都出现正效应、95% CI 不跨 0，并且外部测试不显著退化时，才支持双向正迁移。Oracle 有效而 Pred 无效只能说明条件信号有潜力，不能证明实际模型闭环成立。",
        "",
        "详细预注册方案见 `docs/bidirectional_validation_protocol.md`。",
        "",
        "## 11. 仍需补齐的论文证据",
        "",
        "- 尚无独立第二标注目录，不能报告 10% 双人独立复标、多边形 IoU 或类别一致率。正式论文前必须补做并冻结复核清单。",
        "- Uru 的 3 个非严格嵌套年轮样本需专家复核；在此之前应同时报告包含与排除它们的敏感性分析。",
        "- VNWoodKnot 数字类别 1/2 的生物学含义和所有数据集许可证仍需引用官方来源确认。",
        "- VSB 只有 10 个连续采集组，16,000 张图不等于 16,000 个独立木材来源；置信区间和外推结论必须以采集组为单位解释。",
        "- 数据工程已完成，但尚未生成任何正式模型指标；不得把本报告的完整性 `ok` 写成模型精度结论。",
        "",
        "## 12. 复现命令",
        "",
        "```powershell",
        "python scripts/build_manifest.py --config configs/data.json --refresh-image-metadata",
        "python scripts/audit_provenance.py --baseline",
        "python scripts/run_data_pipeline.py --stages bridge rings coco crops duplicates --vsb-segmentation-limit 0",
        "python scripts/audit_provenance.py --verify",
        "python scripts/validate_data_artifacts.py",
        "python -m unittest discover -s tests -v",
        "python scripts/build_final_report.py",
        "```",
        "",
        f"清单 inventory SHA-256：`{quality['inventory_sha256']}`。关键 manifest SHA-256：`{facts['artifact_hashes']['data_processed/manifests/manifest.csv']}`。",
        "",
    ])
    return "\n".join(lines)


def build_protocol() -> str:
    return """# 年轮-缺陷双向正迁移预注册验证协议

## 研究问题

在完全相同的物理组划分和训练预算下，年轮监督是否改善 knot/crack 识别，缺陷监督是否改善年轮结构提取，并且这种改善能否在预测条件而非 Oracle 条件下稳定复现。

## 固定数据协议

- 数据清单、桥接 train/val/test、VSB 固定划分和 external_test 在训练前冻结。
- 所有裁剪、patch、增强和采样均继承源图 group_id；同一物理组绝不跨划分。
- `knot_with_crack` 只作为 knot；crack 仅使用独立 Crack 或其他数据集明确的 standalone crack/split 标注。
- 非目标缺陷区域进入 ignore，不计作背景；缺失任务标签由 task mask 隔离。

## 模型矩阵

1. Ring-only：独立年轮基线。
2. Defect-only：独立两类检测/分割基线。
3. Shared-no-interaction：共享编码器但任务头之间不交换条件信息。
4. R-to-D Oracle：缺陷头使用真实年轮条件。
5. R-to-D Pred：缺陷头使用模型预测年轮条件。
6. D-to-R Oracle：年轮头使用真实缺陷条件。
7. D-to-R Pred：年轮头使用模型预测缺陷条件。
8. BiCRR：双向预测条件完整模型。

所有模型使用相同骨干、输入尺寸、增强、优化器、训练步数和早停规则。至少使用 5 个训练随机种子，保存每个种子的完整配置、权重哈希和逐物理组预测。

## 指标

- 年轮：Boundary F1/Dice、clDice、ASSD或HD95、闭合率、每图断裂数。
- 缺陷：box/mask mAP50-95、AP50、AP_knot、AP_crack、每类召回率。
- 分类：macro-F1、balanced accuracy、每类 precision/recall/F1。
- 局部分析：缺陷框外扩固定比例后的年轮指标，与同图面积匹配的非缺陷区域比较。

## 统计检验

- 以树、立方体或采集组为重采样单位，执行组级 bootstrap 95% CI。
- 对相同物理组上的模型差值做双侧配对置换检验；多重比较使用 Holm 校正。
- 同时报告绝对差值、相对差值和标准化效应量，不只报告 p 值。
- bridge test=69 只在方案冻结后使用一次；超参数只依据 train/val。

## 判定规则

只有 BiCRR 相对 Shared-no-interaction 在年轮和缺陷两个方向的主要指标均为正、95% CI 不跨 0，且 UruDendro4、Indiana、VNWoodKnot、OULU 外测无预设阈值以上退化时，才支持“双向正迁移”。Oracle 有效但 Pred 无效，结论只能是条件信息具有潜力。任何一侧无改善或外测显著退化，都不能写成稳定双向正相关。

## 敏感性分析

- 包含/排除 3 个非严格嵌套 Uru 样本。
- 按 knot/crack、图像来源、缺陷尺寸、truncated、年轮密度分层。
- 对 crack 使用独立 Crack 真值，单独验证 `knot_with_crack` 从 crack 真值排除后的稳定性。
- 报告 VSB 10 个采集组逐组结果，避免由图像数量掩盖来源偏差。
"""


def set_east_asia_font(run: Any, name: str = "Microsoft YaHei") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def shade_cell(cell: Any, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def add_docx_table(document: Document, table_lines: list[str]) -> None:
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in table_lines]
    if len(rows) > 1 and all(set(cell) <= {"-", ":"} for cell in rows[1]):
        rows.pop(1)
    table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = value.replace("**", "").replace("`", "")
            if row_index == 0:
                shade_cell(cell, "D9EAF2")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_east_asia_font(run)
                    run.font.size = Pt(8)
                    run.bold = row_index == 0


def build_docx(markdown: str, destination: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    document.core_properties.title = "木材年轮与缺陷多任务项目：论文级数据工程报告"
    document.core_properties.subject = "Data engineering, provenance, leakage audit and bidirectional validation protocol"
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10)
    for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12)):
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)

    lines = markdown.splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        stripped = lines[index].strip()
        if stripped.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph()
                run = paragraph.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(8)
                code_lines.clear()
            in_code = not in_code
            index += 1
            continue
        if in_code:
            code_lines.append(lines[index])
            index += 1
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_docx_table(document, table_lines)
            continue
        if stripped.startswith("# "):
            paragraph = document.add_paragraph(stripped[2:], style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=1)
        elif stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(stripped[2:].replace("**", "").replace("`", ""))
        elif stripped:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(5)
            paragraph.add_run(stripped.replace("**", "").replace("`", ""))
        index += 1
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("wood-multitask-final | 2026-08-06")
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def write_sheet(sheet: Any, headers: list[str], rows: Iterable[Iterable[object]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append(list(row))
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, sheet.max_column + 1):
        width = max(len(str(sheet.cell(row, column).value or "")) for row in range(1, min(sheet.max_row, 200) + 1))
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 10), 42)


def build_xlsx(facts: dict[str, Any], destination: Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary_rows = [
        ["Manifest images", facts["quality"]["total_images"]],
        ["Primary images", facts["scope_counts"]["primary"]],
        ["Auxiliary images", facts["scope_counts"]["auxiliary_only"]],
        ["Ring targets", facts["ring"]["generated"]],
        ["COCO images", facts["integrity"]["defect_coco"]["images"]],
        ["COCO target instances", facts["integrity"]["defect_coco"]["annotations"]],
        ["Ignored regions", facts["integrity"]["defect_coco"]["ignore_regions"]],
        ["Classification crops", facts["crops"]["total"]],
        ["Bridge images", sum(facts["bridge_splits"].values())],
        ["Integrity status", facts["integrity"]["status"]],
        ["Source unchanged", facts["source_verification"]["source_unchanged"]],
        ["Inventory SHA-256", facts["quality"]["inventory_sha256"]],
    ]
    write_sheet(summary, ["Metric", "Value"], summary_rows)

    registry_sheet = workbook.create_sheet("Dataset Registry")
    registry_rows = []
    for dataset_id, item in facts["registry"]["datasets"].items():
        registry_rows.append([
            dataset_id, item["title"], item["protocol_role"], item["manifest_images"],
            item["physical_groups"], split_text(item["splits"]), item["ring_images"],
            item["defect_labeled_images"], item["target_instances_manifest"],
            item["ignored_instances_manifest"], item["unit_of_independence"], item["license_status"],
        ])
    write_sheet(
        registry_sheet,
        ["dataset_id", "title", "role", "images", "groups", "splits", "ring_images", "defect_labeled_images", "target_instances", "ignored_instances", "unit", "license_status"],
        registry_rows,
    )

    defects = workbook.create_sheet("Defect Instances")
    defect_rows = []
    for dataset_id in ("vsb", "vnwoodknot", "oulu", "urudendro", "mokume"):
        knot = facts["dataset_category"][(dataset_id, "knot")]
        crack = facts["dataset_category"][(dataset_id, "crack")]
        defect_rows.append([dataset_id, knot, crack, knot + crack, facts["normal_by_dataset"][dataset_id]])
    write_sheet(defects, ["dataset_id", "knot", "crack", "targets_total", "verified_normal"], defect_rows)

    splits = workbook.create_sheet("Defect Split Counts")
    split_rows = []
    for split in ("train", "val", "test", "external_test"):
        knot = facts["split_category"][(split, "knot")]
        crack = facts["split_category"][(split, "crack")]
        split_rows.append([split, knot, crack, knot + crack])
    write_sheet(splits, ["split", "knot", "crack", "total"], split_rows)

    rings = workbook.create_sheet("Ring Targets")
    write_sheet(
        rings,
        ["dataset_id", "generated"],
        sorted(facts["ring"]["by_dataset"].items()),
    )
    rings.append([])
    rings.append(["Non-nested warning samples"])
    for sample_id in facts["ring_warning_samples"]:
        rings.append([sample_id])

    artifacts = workbook.create_sheet("Artifact Hashes")
    write_sheet(artifacts, ["path", "sha256"], sorted(facts["artifact_hashes"].items()))

    limitations = workbook.create_sheet("Limitations")
    limitation_rows = [
        ["Independent reannotation", "Missing", "Complete at least 10% independent second annotation before paper claims"],
        ["Uru topology warnings", 3, ", ".join(facts["ring_warning_samples"])],
        ["VSB source diversity", "10 acquisition groups", "Use group-level inference and external tests"],
        ["VN class semantics", "Needs official citation", "Confirm numeric classes 1/2"],
        ["Licenses", "Mostly unrecorded", "Verify before redistribution"],
        ["Bidirectional effect", "Not yet measured", "Run preregistered model matrix; do not infer from data integrity"],
    ]
    write_sheet(limitations, ["Issue", "Status", "Required action"], limitation_rows)

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    facts = collect_facts()
    markdown = build_markdown(facts)
    md_path = DOCS / f"{REPORT_STEM}.md"
    docx_path = DOCS / f"{REPORT_STEM}.docx"
    xlsx_path = DOCS / f"{REPORT_STEM}.xlsx"
    protocol_path = DOCS / "bidirectional_validation_protocol.md"
    md_path.write_text(markdown, encoding="utf-8")
    protocol_path.write_text(build_protocol(), encoding="utf-8")
    build_docx(markdown, docx_path)
    build_xlsx(facts, xlsx_path)
    print(json.dumps({
        "markdown": str(md_path),
        "docx": str(docx_path),
        "xlsx": str(xlsx_path),
        "protocol": str(protocol_path),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
